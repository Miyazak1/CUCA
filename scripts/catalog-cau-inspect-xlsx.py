import json
from pathlib import Path
from openpyxl import load_workbook

root = Path("work/catalog-official/cau-program-attachments-batch-01/raw")
result = []
for path in sorted(root.glob("*.xlsx")):
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        populated = []
        for index, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = [value for value in row]
            if any(value not in (None, "") for value in values):
                populated.append({"row": index, "values": values})
        sheets.append({"title": ws.title, "maxRow": ws.max_row, "maxColumn": ws.max_column, "populatedRows": len(populated), "sample": populated[:5] + populated[-3:]})
    result.append({"file": path.name, "sheets": sheets})
print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
