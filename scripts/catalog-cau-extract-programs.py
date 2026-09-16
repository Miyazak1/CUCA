import json
import re
from pathlib import Path
from openpyxl import load_workbook

root = Path("work/catalog-official/cau-program-attachments-batch-01")
manifest = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
sources = {row["id"]: row for row in manifest["sources"]}

def workbook(source_id):
    path = root / sources[source_id]["artifactPath"]
    return load_workbook(path, read_only=True, data_only=True)

def english(value):
    if value is None:
        return ""
    text = str(value).replace("\u00a0", " ")
    parts = re.findall(r"[A-Za-z][A-Za-z0-9&(),.'’/\- ]*", text)
    return re.sub(r"\s+", " ", " ".join(parts)).strip(" /-")

def duration(value):
    match = re.search(r"(\d+(?:\.\d+)?)", str(value or ""))
    return float(match.group(1)) if match else None

programs = []
undergraduate = workbook("cau-undergraduate-program-catalog-2026").worksheets[0]
college = ""
for row_number, row in enumerate(undergraduate.iter_rows(min_row=3, values_only=True), start=3):
    if row_number == undergraduate.max_row:
        break
    if row[0]:
        college = english(row[0])
    name = english(str(row[1] or "").splitlines()[0]) or english(row[1])
    if not name:
        continue
    subject_text = english(row[3])
    subjects = [part.strip() for part in subject_text.split(",") if part.strip()]
    programs.append({
        "sourceId": "cau-undergraduate-program-catalog-2026",
        "sourceRow": row_number,
        "collegeEn": college,
        "nameEn": name,
        "degreeLevel": "Undergraduate",
        "teachingLanguage": "Chinese",
        "cscaSubjects": subjects,
    })

graduate = workbook("cau-graduate-program-catalog-2026").worksheets[0]
college = ""
for row_number, row in enumerate(graduate.iter_rows(min_row=3, values_only=True), start=3):
    if row[0]:
        college = english(row[0])
    first_line = str(row[1] or "").splitlines()[0]
    name = english(first_line) or english(row[1])
    if not name:
        continue
    languages = []
    language_text = english(row[2]).lower()
    if "english" in language_text:
        languages.append("English")
    if "chinese" in language_text:
        languages.append("Chinese")
    for degree, cell in (("Master", row[3]), ("Doctoral", row[4])):
        years = duration(cell)
        if years is None:
            continue
        for language in languages:
            programs.append({
                "sourceId": "cau-graduate-program-catalog-2026",
                "sourceRow": row_number,
                "collegeEn": college,
                "nameEn": name,
                "degreeLevel": degree,
                "teachingLanguage": language,
                "durationYears": years,
                "cscaSubjects": [],
            })

identities = [(p["degreeLevel"], p["nameEn"], p["teachingLanguage"], p["collegeEn"]) for p in programs]
duplicates = sorted({"|".join(key) for key in identities if identities.count(key) > 1})
counts = {}
for program in programs:
    key = f'{program["degreeLevel"]}:{program["teachingLanguage"]}'
    counts[key] = counts.get(key, 0) + 1
result = {"programCount": len(programs), "counts": counts, "duplicateIdentities": duplicates, "programs": programs}
(root / "parsed-programs.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({key: result[key] for key in ("programCount", "counts", "duplicateIdentities")}, ensure_ascii=False, indent=2))
