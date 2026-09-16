from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "work/catalog-official/whu-program-docx-browser-20260914/raw"
OUTPUT = ROOT / "work/catalog-official/whu-program-docx-browser-20260914/parsed-programs.json"

SOURCES = {
    "whu-undergraduate-programs-chinese-2026": {
        "pattern": "whu-undergraduate-programs-chinese-2026.*.docx",
        "degree": "Undergraduate",
        "language": "Chinese",
        "columns": (0, 1, 2),
    },
    "whu-undergraduate-programs-english-2026": {
        "pattern": "whu-undergraduate-programs-english-2026.*.docx",
        "degree": "Undergraduate",
        "language": "English",
        "columns": (0, 1, 2),
    },
    "whu-master-programs-chinese-2026": {
        "pattern": "whu-master-programs-chinese-2026.*.docx",
        "degree": "Master",
        "language": "Chinese",
        "columns": (0, 1, 2),
    },
    "whu-doctoral-programs-chinese-2026": {
        "pattern": "whu-doctoral-programs-chinese-2026.*.docx",
        "degree": "Doctoral",
        "language": "Chinese",
        "columns": (0, 1, 2),
    },
    "whu-graduate-programs-english-2026": {
        "pattern": "whu-graduate-programs-english-2026.*.docx",
        "degree": None,
        "language": "English",
        "columns": (0, 1, 2),
    },
}


def normalize(value: str) -> str:
    value = value.replace("\xa0", " ").replace("\u3000", " ")
    value = value.replace("（修改）", "").replace("(修改)", "")
    return re.sub(r"\s+", " ", value).strip(" /\t\r\n")


def split_bilingual(value: str, chinese_first: bool) -> tuple[str | None, str | None]:
    value = normalize(value)
    if not value:
        return None, None
    parts = [normalize(part) for part in re.split(r"\s*/\s*", value) if normalize(part)]
    if len(parts) >= 2:
        chinese = next((part for part in parts if re.search(r"[\u3400-\u9fff]", part)), None)
        english = next((part for part in parts if re.search(r"[A-Za-z]", part)), None)
        if chinese and english and chinese != english:
            return (chinese, english) if chinese_first else (chinese, english)
    match = re.search(r"[\u3400-\u9fff]", value)
    if match:
        left, right = normalize(value[: match.start()]), normalize(value[match.start() :])
        if chinese_first:
            chinese = normalize(re.match(r"^[\u3400-\u9fff（）()·\-\s]+", value).group(0)) if re.match(r"^[\u3400-\u9fff（）()·\-\s]+", value) else value
            english = normalize(value[len(chinese) :])
        else:
            english, chinese = left, right
        return chinese or None, english or None
    return (value, None) if chinese_first else (None, value)


def duration_years(value: str) -> int | None:
    match = re.search(r"(?:^|\D)([2-6])\s*(?:年|Years?|years?)", normalize(value), re.I)
    return int(match.group(1)) if match else None


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def extract_source(source_id: str, config: dict[str, object]) -> tuple[dict[str, object], list[dict[str, object]]]:
    matches = list(RAW_DIR.glob(str(config["pattern"])))
    if len(matches) != 1:
        raise RuntimeError(f"Expected one DOCX for {source_id}, found {len(matches)}")
    path = matches[0]
    document = Document(path)
    if len(document.tables) != 1:
        raise RuntimeError(f"Expected one table in {path.name}, found {len(document.tables)}")
    table = document.tables[0]
    rows: list[dict[str, object]] = []
    last_department_raw = ""
    catalog_group = ""
    for index, row in enumerate(table.rows):
        values = [normalize(cell.text.replace("\n", " / ")) for cell in row.cells]
        department_col, program_col, duration_col = config["columns"]
        department_raw = values[int(department_col)]
        program_raw = values[int(program_col)]
        duration_raw = values[int(duration_col)]
        years = duration_years(duration_raw)
        if years is None:
            group_match = re.search(r"Liberal Arts|Science and Engineering|Medicine|Interdisciplinary", " ".join(values), re.I)
            if group_match:
                catalog_group = group_match.group(0).title()
        if department_raw and years:
            last_department_raw = department_raw
        if not program_raw or years is None:
            continue
        degree = config["degree"]
        if source_id == "whu-graduate-programs-english-2026":
            degree = "Doctoral" if re.search(r"Doctor", department_raw, re.I) else "Master"
        chinese_first = config["language"] == "Chinese"
        name_zh, name_en = split_bilingual(program_raw, chinese_first)
        department_zh, department_en = split_bilingual(last_department_raw or department_raw, chinese_first)
        if not name_en:
            raise RuntimeError(f"Missing English program name in {source_id} row {index}: {values}")
        note = values[3] if len(values) > 3 else ""
        professional = bool(re.search(r"professional degree|专业学位", f"{department_raw} {program_raw} {note}", re.I))
        rows.append(
            {
                "sourceId": source_id,
                "sourceFile": path.name,
                "sourceRow": index,
                "degreeLevel": degree,
                "teachingLanguage": config["language"],
                "nameEn": name_en,
                **({"nameZh": name_zh} if name_zh else {}),
                **({"departmentEn": department_en} if department_en else {}),
                **({"departmentZh": department_zh} if department_zh else {}),
                **({"catalogGroup": catalog_group} if catalog_group else {}),
                "durationYears": years,
                "degreeType": "professional" if professional else "academic-or-unspecified",
                "raw": values,
            }
        )
    return {
        "sourceId": source_id,
        "file": path.name,
        "sha256": sha256(path),
        "tableRows": len(table.rows),
        "extractedRows": len(rows),
    }, rows


def main() -> None:
    source_summaries: list[dict[str, object]] = []
    programs: list[dict[str, object]] = []
    for source_id, config in SOURCES.items():
        summary, rows = extract_source(source_id, config)
        source_summaries.append(summary)
        programs.extend(rows)
    keys: dict[tuple[str, str, str, str, str], list[int]] = {}
    for index, row in enumerate(programs):
        key = (
            str(row["degreeLevel"]),
            str(row["teachingLanguage"]),
            str(row["nameEn"]).casefold(),
            str(row.get("departmentEn", "")).casefold(),
            str(row["degreeType"]),
        )
        keys.setdefault(key, []).append(index)
    duplicates = [{"key": list(key), "indexes": indexes} for key, indexes in keys.items() if len(indexes) > 1]
    payload = {
        "version": 1,
        "sourceManifest": "work/catalog-official/whu-program-docx-browser-20260914/manifest.json",
        "summary": {
            "sources": source_summaries,
            "programRoutes": len(programs),
            "byDegreeLanguage": {
                f"{degree}:{language}": sum(
                    1 for row in programs if row["degreeLevel"] == degree and row["teachingLanguage"] == language
                )
                for degree, language in sorted({(str(row["degreeLevel"]), str(row["teachingLanguage"])) for row in programs})
            },
            "duplicateRouteKeys": len(duplicates),
        },
        "duplicates": duplicates,
        "programs": programs,
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload["summary"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"WHU DOCX extraction failed: {error}", file=sys.stderr)
        raise
