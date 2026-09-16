from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_HASHES = {
    "csu-chinese-undergraduate-programs-2026.xlsx": "a0e72fcd0cae22ac4ca67658a470b877333aba7a45242acf2682bbcf9cbc7ee8",
    "csu-english-undergraduate-programs-2026.xlsx": "72e5830c78f875cf81236f981fe7a3d20a7c3bb03d59548da45cc3403e4d3717",
    "csu-graduate-programs-2026.xlsx": "a83d3d86ca3c9324ed31b1c73c651297b47078e9159413784d38b1961017d979",
}


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\xa0", " ")).strip()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def split_school(value: str) -> tuple[str, str]:
    match = re.search(r"[A-Za-z]", value)
    if not match:
        raise ValueError(f"No English school name in undergraduate school cell: {value!r}")
    return clean(value[: match.start()]), clean(value[match.start() :])


def parse_undergraduate(path: Path, expected_language: str) -> list[dict[str, object]]:
    sheet = load_workbook(path, data_only=True, read_only=True).active
    routes: list[dict[str, object]] = []
    school_zh = school_en = ""
    for row in sheet.iter_rows(min_row=3, values_only=True):
        cells = [clean(value) for value in row]
        if not cells[2]:
            continue
        if cells[0]:
            school_zh, school_en = split_school(cells[0])
        if not school_en:
            raise ValueError(f"Missing carried-forward school for {cells[2]}")
        language = cells[4]
        if language != expected_language:
            raise ValueError(f"Unexpected undergraduate language {language!r} for {cells[2]}")
        years_match = re.search(r"\d+", cells[6])
        if not years_match:
            raise ValueError(f"Missing duration for {cells[2]}")
        routes.append(
            {
                "degree": "Undergraduate",
                "schoolZh": school_zh,
                "schoolEn": school_en,
                "nameZh": cells[1],
                "nameEn": cells[2],
                "teachingLanguage": language,
                "durationYears": int(years_match.group()),
                "cscaRequirement": cells[8],
                "researchFieldsZh": [],
                "researchFieldsEn": [],
                "notesZh": [],
                "notesEn": [],
                "sourceWorkbook": path.name,
                "sourceSheet": sheet.title,
            }
        )
    return routes


def parse_graduate(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    routes: list[dict[str, object]] = []
    for sheet in workbook.worksheets:
        degree = "Master" if "Master" in sheet.title else "Doctoral"
        grouped: dict[tuple[str, ...], dict[str, object]] = {}
        for row in sheet.iter_rows(min_row=3, values_only=True):
            cells = [clean(value) for value in row]
            if not cells[1] or not cells[3] or not cells[7]:
                continue
            years_match = re.search(r"\d+", cells[9])
            if not years_match:
                raise ValueError(f"Missing duration for {cells[3]} in {sheet.title}")
            # Chinese degree names distinguish academic and professional routes
            # that CSU sometimes publishes under the same English major label.
            key = (degree, cells[1], cells[2], cells[3], cells[7])
            route = grouped.setdefault(
                key,
                {
                    "degree": degree,
                    "schoolZh": cells[0],
                    "schoolEn": cells[1],
                    "nameZh": cells[2],
                    "nameEn": cells[3],
                    "teachingLanguage": cells[7],
                    "durationYears": int(years_match.group()),
                    "cscaRequirement": None,
                    "researchFieldsZh": [],
                    "researchFieldsEn": [],
                    "notesZh": [],
                    "notesEn": [],
                    "sourceWorkbook": path.name,
                    "sourceSheet": sheet.title,
                },
            )
            for target, value in (
                ("researchFieldsZh", cells[4]),
                ("researchFieldsEn", cells[5]),
                ("notesZh", cells[10]),
                ("notesEn", cells[11]),
            ):
                values = route[target]
                if value and value not in values:
                    values.append(value)
        routes.extend(grouped.values())
    return routes


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", default="work/catalog-official/csu-complete-batch-01/manual")
    parser.add_argument("--output", default="work/catalog-official/csu-complete-batch-01/manual/program-routes.json")
    args = parser.parse_args()
    input_dir = Path(args.input_dir)
    output = Path(args.output)

    for name, expected in EXPECTED_HASHES.items():
        actual = sha256(input_dir / name)
        if actual != expected:
            raise ValueError(f"Official attachment hash mismatch for {name}: {actual}")

    routes = [
        *parse_undergraduate(input_dir / "csu-chinese-undergraduate-programs-2026.xlsx", "Chinese"),
        *parse_undergraduate(input_dir / "csu-english-undergraduate-programs-2026.xlsx", "English"),
        *parse_graduate(input_dir / "csu-graduate-programs-2026.xlsx"),
    ]
    counts = Counter(route["degree"] for route in routes)
    expected_counts = {"Undergraduate": 84, "Master": 166, "Doctoral": 120}
    if dict(counts) != expected_counts:
        raise ValueError(f"Unexpected route counts: {dict(counts)}")
    languages = defaultdict(Counter)
    for route in routes:
        languages[route["degree"]][route["teachingLanguage"]] += 1
    route_keys = {
        (route["degree"], route["schoolEn"], route["nameZh"], route["nameEn"], route["teachingLanguage"])
        for route in routes
    }
    if len(route_keys) != len(routes):
        raise ValueError("Program-route aggregation still contains duplicate keys")

    payload = {
        "version": 1,
        "sourceHashes": EXPECTED_HASHES,
        "counts": {
            "routes": len(routes),
            "byDegree": dict(counts),
            "byLanguageAndDegree": {degree: dict(values) for degree, values in languages.items()},
        },
        "aggregationRule": "One route per degree + school + Chinese major + English major + teaching language; research-field rows are retained as route detail.",
        "routes": routes,
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"ok": True, "output": str(output), **payload["counts"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
